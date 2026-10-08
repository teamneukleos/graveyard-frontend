from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "1A1A1A")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(255, 255, 255)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
        if r_idx % 2 == 1:
            for c in cells:
                set_cell_shading(c, "F4F4F4")
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Graveyard — UI QA Test Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph()
    run = p.add_run(
        "Thorough end-to-end QA checklist covering every user type and major flow."
    )
    run.italic = True

    meta = doc.add_paragraph()
    meta.add_run("Version: ").bold = True
    meta.add_run("1.0  |  ")
    meta.add_run("Product: ").bold = True
    meta.add_run("Graveyard  |  ")
    meta.add_run("Mark each case: ").bold = True
    meta.add_run("Pass / Fail / Blocked / N/A + notes")

    doc.add_paragraph(
        "Suggested environments: Staging first, then local (API :3000, App :3001)."
    )

    doc.add_heading("0. Accounts to prepare before testing", level=1)
    doc.add_paragraph(
        "Create these accounts (or reuse existing) before starting the suite:"
    )
    add_table(
        doc,
        ["#", "Account", "How to create", "Purpose"],
        [
            ["1", "Seed Admin", "admin@graveyard.local / ChangeMeAdmin1!", "Full admin flows"],
            ["2", "Email Creator A", "Register as Creator + verify email", "Submit, publish, awards"],
            ["3", "Email Creator B", "Second creator, verified", "Like/follow/cross-user"],
            ["4", "Email Agency", "Register as Agency with agency name + verify", "Agency portal & profile"],
            ["5", "Google Agency", "Register → Agency → Continue with Google", "Onboarding path"],
            ["6", "Judge", "Create from /admin/judges, assign to cycle", "Scoring queue"],
        ],
    )

    doc.add_heading("1. Guest (not logged in)", level=1)
    doc.add_paragraph(
        "Surfaces: /, /showcase, /showcase/[slug], /categories, /leaderboards, "
        "/creators/[id], /agencies/[slug], /events, legal pages."
    )
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["G1", "Browse home", "Open /, filter by category/status/search", "Feed loads; filters work", "", ""],
            ["G2", "Showcase list", "Open /showcase, open a piece", "Detail, gallery, concept show", "", ""],
            ["G3", "Like gated", "Click Like while logged out", "Auth modal (Sign in / Create account / Google)", "", ""],
            ["G4", "Follow gated", "On a profile, click Follow", "Same auth modal", "", ""],
            ["G5", "Return after auth", "From modal → login/register with next=", "Returns to same page after auth", "", ""],
            ["G6", "Share links", "Copy profile/project link; open in private window", "Public page loads", "", ""],
            ["G7", "Leaderboards", "Open /leaderboards (+ creators/agencies)", "Boards render", "", ""],
            ["G8", "Protected routes", "Visit /portal, /admin, /judge, /settings", "Redirect to /login?next=…", "", ""],
            ["G9", "Events browse", "Open /events", "Events list loads", "", ""],
            ["G10", "Legal pages", "Open /terms, /privacy, /cookies", "Pages render", "", ""],
        ],
    )

    doc.add_heading("2. Creator (email)", level=1)
    doc.add_heading("2A. Register & verify", level=2)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["C1", "Register", "/register → Creator → name, email, password (≥8)", "“Verify your email”; not logged in", "", ""],
            ["C2", "Unverified login", "Login before verifying", "Can enter app but gated actions blocked", "", ""],
            ["C3", "Unverified submit", "Open /portal/submit", "Banner / blocked submit until verified", "", ""],
            ["C4", "Unverified like", "Like someone’s work", "Error: verify email", "", ""],
            ["C5", "Verify email", "Open verify link → /verify-email?token=", "Success; full access after login", "", ""],
            ["C6", "Resend verify", "Use resend if available", "New link works", "", ""],
        ],
    )

    doc.add_heading("2B. Submit & publish", level=2)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["C7", "Create draft", "/portal/submit → fill form → save", "Draft at /portal/submissions/[id]", "", ""],
            ["C8", "Upload assets", "Upload image(s)", "Image displays; no broken localhost URL", "", ""],
            ["C9", "Edit draft", "Change title/concept/assets", "Changes save", "", ""],
            ["C10", "Publish validation", "Publish without rights / too-short copy", "Blocked with clear error", "", ""],
            ["C11", "Publish", "Attest rights → Publish", "Published; visible at /showcase/[slug]", "", ""],
            ["C12", "Post-publish edit", "Try editing fields", "Read-only; award entry still possible if open", "", ""],
            ["C13", "Portal list", "Open /portal", "Card shows; share/profile links work", "", ""],
            ["C14", "Public profile", "Open /creators/{id}", "Work grid, share, follow (from other user)", "", ""],
        ],
    )

    doc.add_heading("2C. Social", level=2)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["C15", "Like others", "Like another creator’s piece", "Count updates; Liked state", "", ""],
            ["C16", "Unlike", "Toggle like off", "Count drops", "", ""],
            ["C17", "Self-like", "Like own piece", "Rejected", "", ""],
            ["C18", "Follow", "Follow another creator/agency", "Following state updates", "", ""],
            ["C19", "Self-follow", "Follow self", "Rejected", "", ""],
            ["C20", "Share", "Copy profile + project links", "Open correctly while logged out", "", ""],
        ],
    )

    doc.add_heading("2D. Awards (creator side)", level=2)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["C21", "Enter cycle", "Published piece → Enter open cycle", "Entry succeeds", "", ""],
            ["C22", "Double enter", "Enter same cycle again", "Rejected / already entered", "", ""],
            ["C23", "Withdraw", "Withdraw while cycle open", "Removed from cycle", "", ""],
            ["C24", "Closed cycle", "Try enter closed cycle", "Rejected", "", ""],
        ],
    )

    doc.add_heading("2E. Account", level=2)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["C25", "Settings", "/settings → name, bio, avatar", "Saves; avatar shows", "", ""],
            ["C26", "Password reset", "Forgot → email → reset → login", "Works", "", ""],
            ["C27", "Logout", "Log out", "Becomes guest; protected routes blocked", "", ""],
        ],
    )

    doc.add_heading("3. Agency (email)", level=1)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["A1", "Register", "/register → Agency → agency name + email + password", "Verify flow same as creator", "", ""],
            ["A2", "Login home", "After verify → /portal", "No forced onboarding", "", ""],
            ["A3", "Submit/publish", "Same as creator", "Public under /agencies/{id} (or name)", "", ""],
            ["A4", "Agency profile", "Open public agency page", "Work, team, share, follow", "", ""],
            ["A5", "Like weight", "Agency likes a piece", "Score increases more than a creator like", "", ""],
            ["A6", "Settings", "Update agency name / bio / avatar", "Public profile updates", "", ""],
            ["A7", "Incomplete agency", "Agency without name tries vote/submit", "Blocked until onboarding complete", "", ""],
        ],
    )

    doc.add_heading("4. Agency (Google) — onboarding", level=1)
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["GA1", "Google as agency", "Register → Agency → Continue with Google", "Account created (email verified)", "", ""],
            ["GA2", "Forced onboarding", "Land on /onboarding/agency", "Cannot usefully skip", "", ""],
            ["GA3", "Blocked actions", "Before name: like / follow / submit", "Forbidden or redirected", "", ""],
            ["GA4", "Complete name", "Submit agency name", "Goes to /portal; flows unlock", "", ""],
            ["GA5", "Google without role", "Google from /login with no role", "New user becomes Creator (expected)", "", ""],
        ],
    )

    doc.add_heading("5. Judge", level=1)
    doc.add_paragraph(
        "Setup: Admin creates judge at /admin/judges, then assigns them to a cycle at /admin/cycles."
    )
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["J1", "Login home", "Judge login → /judge", "Queue for active cycle", "", ""],
            ["J2", "Open piece", "Open queue item", "Detail + assets + score form", "", ""],
            ["J3", "Score", "Submit score in valid range", "Saved; marked scored", "", ""],
            ["J4", "Re-score", "Change score", "Updates", "", ""],
            ["J5", "Not assigned", "Judge on cycle they are not on", "No access / forbidden", "", ""],
            ["J6", "Own work", "If assigned own piece", "Cannot score", "", ""],
            ["J7", "Outside JUDGING", "Score when upcoming/closed", "Rejected", "", ""],
            ["J8", "Wrong routes", "Visit /admin or /portal", "Redirect away", "", ""],
            ["J9", "Settings", "Update profile", "Works", "", ""],
        ],
    )

    doc.add_heading("6. Admin / Super Admin", level=1)
    doc.add_paragraph(
        "Seed: admin@graveyard.local / ChangeMeAdmin1! → should land on /admin."
    )
    add_table(
        doc,
        ["ID", "Flow", "Steps", "Pass if", "Result", "Notes"],
        [
            ["AD1", "Login home", "Login as admin", "Lands on /admin (not portal)", "", ""],
            ["AD2", "Submissions table", "View / filter / open", "Lists work", "", ""],
            ["AD3", "Bulk publish", "Publish selected drafts", "Statuses update", "", ""],
            ["AD4", "Shortlist/winners", "Use admin status/winner paths", "Showcase reflects", "", ""],
            ["AD5", "Categories", "/admin/categories CRUD", "Appear in submit + browse", "", ""],
            ["AD6", "Judges", "/admin/judges create/list", "Judge can log in", "", ""],
            ["AD7", "Create cycle", "/admin/cycles → Create", "Appears in table", "", ""],
            ["AD8", "View/Edit cycle", "View + Edit modal", "Dates/status/judges save", "", ""],
            ["AD9", "Assign judge", "Assign from edit", "Judge sees queue", "", ""],
            ["AD10", "Status flow", "UPCOMING → JUDGING → RESULTS_PUBLISHED → CLOSED", "Illegal jumps blocked", "", ""],
            ["AD11", "Events", "/admin/events create/edit", "Public /events shows them", "", ""],
            ["AD12", "Analytics", "/admin/analytics", "Loads without error", "", ""],
            ["AD13", "Portal redirect", "Visit /portal as admin", "Redirects to /admin", "", ""],
        ],
    )

    doc.add_heading("7. End-to-end award story (multi-role)", level=1)
    doc.add_paragraph("Run once as a full regression:")
    for step in [
        "Admin creates cycle (UPCOMING) and assigns Judge",
        "Creator publishes a project and enters the cycle",
        "Judge scores the entry",
        "Admin advances cycle / publishes results",
        "Guest sees the piece on /showcase with correct badges",
        "Other creator likes it; leaderboards update",
    ]:
        doc.add_paragraph(step, style="List Number")

    e2e = doc.add_paragraph()
    e2e.add_run("E2E Result: ").bold = True
    e2e.add_run("Pass / Fail _____________   Notes: ________________________________")

    doc.add_heading("8. Cross-cutting checks", level=1)
    add_table(
        doc,
        ["ID", "Area", "Check", "Result", "Notes"],
        [
            ["X1", "Auth modal", "Like/Follow logged out → modal, not silent fail", "", ""],
            ["X2", "Images", "Upload displays on portal + public (no localhost URLs on staging)", "", ""],
            ["X3", "Share", "Profile + project links work in private window", "", ""],
            ["X4", "Mobile", "Key pages readable; like button visible; modals usable", "", ""],
            ["X5", "Events RSVP", "/events RSVP requires login", "", ""],
            ["X6", "Session", "Refresh keeps login; logout clears access", "", ""],
            ["X7", "Role walls", "Creator cannot open /admin or /judge; judge cannot open /admin", "", ""],
            ["X8", "Like UI", "Like button noticeable on feed cards + showcase", "", ""],
        ],
    )

    doc.add_heading("9. Suggested test order (if time is limited)", level=1)
    for item in [
        "Guest browse + like auth gate",
        "Creator register → verify → submit → publish → public page",
        "Upload/display images (local + staging)",
        "Agency register + Google onboarding",
        "Admin cycle + judge scoring + showcase results",
        "Follow / share / leaderboards / settings / password reset",
        "Negative cases (self-like, unverified vote, closed cycle, wrong-role URLs)",
    ]:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("10. Role × surface matrix", level=1)
    add_table(
        doc,
        ["Surface", "Guest", "Creator", "Agency", "Judge", "Admin"],
        [
            ["Explore / showcase / profiles", "Read", "Read", "Read", "Read", "Read"],
            ["Like / follow", "Auth gate", "Yes*", "Yes*", "Yes*", "Yes*"],
            ["/portal submit", "—", "Yes*", "Yes*", "—", "Redirect /admin"],
            ["/judge", "—", "—", "—", "Yes", "Yes"],
            ["/admin/*", "—", "—", "—", "—", "Yes"],
            ["/settings", "—", "Yes", "Yes", "Yes", "Yes"],
        ],
    )
    doc.add_paragraph(
        "* Requires verified email; agencies also need completed onboarding."
    )

    doc.add_heading("11. Sign-off", level=1)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Tester name", ""],
            ["Environment (staging/local)", ""],
            ["Build / date", ""],
            ["Browser(s)", ""],
            ["Overall result", "Pass / Fail / Pass with issues"],
            ["Blocking bugs found", ""],
            ["Non-blocking issues", ""],
            ["Sign-off", ""],
        ],
    )

    out = r"c:\Users\hp\Desktop\graveyard\Graveyard-UI-QA-Test-Guide.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
